"""Append the current drug-handoff contrast to the two-slide visual deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from create_adapter_resolver_visual import (
    CHEMBL,
    CHEMBL_PALE,
    DC,
    DC_PALE,
    INK,
    LINE,
    MERGE,
    MERGE_PALE,
    MUTED,
    NAVY,
    RESOLVER,
    WHITE,
    arrow,
    box,
    field,
    provenance_tag,
    text,
)


HERE = Path(__file__).parent
SOURCE = HERE / "adapter_resolver_visual_with_protein.pptx"
OUT = HERE / "adapter_resolver_visual_complete.pptx"
RED = RGBColor(184, 67, 67)
RED_PALE = RGBColor(250, 232, 230)


def source_card(slide, title_value, rows, x, y, fill, accent):
    box(slide, x, y, 1.92, 1.78, fill, accent)
    text(slide, title_value, x + .13, y + .12, 1.66, .42,
         13, accent, True, PP_ALIGN.CENTER)
    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x + .22), Inches(y + .61), Inches(1.48), Inches(.025)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = accent
    divider.line.fill.background()
    for index, (label, value) in enumerate(rows):
        size = 7.8 if label == "ID" else 10.2
        field(slide, label, value, x + .14, y + .75 + index * .29, 1.64, accent, size)


def app_graph_card(slide, x, y):
    box(slide, x, y, 2.00, 3.82, MERGE_PALE, MERGE)
    text(slide, "COMBINED EXPORT", x + .14, y + .15, 1.72, .25,
         10.5, MERGE, True, PP_ALIGN.CENTER)
    text(slide, "drug_nodes_full.tsv", x + .12, y + .45, 1.76, .36,
         13.5, MERGE, True, PP_ALIGN.CENTER)
    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x + .22), Inches(y + .91), Inches(1.56), Inches(.025)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = MERGE
    divider.line.fill.background()
    text(slide, "one already-canonical row", x + .18, y + 1.05, 1.64, .25,
         10.0, MUTED, True, PP_ALIGN.CENTER)
    text(slide, "drug_id", x + .16, y + 1.48, .64, .22, 9.8, MUTED, True)
    text(slide, "IFXDrug:000042", x + .16, y + 1.73, 1.68, .22, 9.5, INK)
    text(slide, "standard_name", x + .16, y + 2.08, 1.10, .22, 9.8, MUTED, True)
    text(slide, "Example Drug", x + .16, y + 2.33, 1.68, .22, 10.2, INK)
    text(slide, "source_namespaces", x + .16, y + 2.68, 1.40, .22, 9.4, MUTED, True)
    text(slide, "DrugCentral | ChEMBL", x + .16, y + 2.93, 1.68, .25, 9.3, INK)
    box(slide, x + .17, y + 3.30, 1.66, .34, WHITE, MERGE)
    text(slide, "No per-field attribution", x + .25, y + 3.36, 1.50, .20,
         9.5, RED, True, PP_ALIGN.CENTER)


def add_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    text(slide, "With drug_nodes_full.tsv, ODIN sees the result—not the original contributions",
         .52, .27, 12.28, .52, 23.5, NAVY, True)
    text(slide,
         "Identity resolution and payload selection both happen upstream; ODIN receives one combined Drug Harmonizer record.",
         .54, .84, 12.10, .32, 13, MUTED)
    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(.55), Inches(1.25), Inches(.92), Inches(.04)
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = RED
    rule.line.fill.background()

    source_card(
        slide,
        "DATA SOURCE A\nDrugCentral",
        [("ID", "DrugCentral:DC123"), ("name", "Example Drug"), ("status", "approved")],
        .45, 1.55, DC_PALE, DC,
    )
    source_card(
        slide,
        "DATA SOURCE B\nChEMBL",
        [("ID", "ChEMBL:CHEMBL456"), ("name", "Example Drug"), ("phase", "4")],
        .45, 4.36, CHEMBL_PALE, CHEMBL,
    )

    box(slide, 2.72, 1.47, 2.72, 4.82, RED_PALE, RED)
    text(slide, "IFX HARMONIZERS", 2.94, 1.69, 2.28, .32,
         16, RED, True, PP_ALIGN.CENTER)
    box(slide, 3.05, 2.22, 2.06, 1.30, WHITE, RED)
    text(slide, "WHO?", 3.22, 2.38, .60, .25, 11.5, RED, True)
    text(slide, "Link the source IDs and\nchoose IFXDrug:000042", 3.84, 2.33, 1.10, .72,
         10.3, INK, False, PP_ALIGN.LEFT)
    box(slide, 3.05, 3.82, 2.06, 1.61, WHITE, RED)
    text(slide, "WHAT?", 3.22, 3.98, .64, .25, 11.5, RED, True)
    text(slide, "Choose exported fields,\ncombine source values,\nand derive isDrug", 3.87, 3.93, 1.07, 1.02,
         10.2, INK, False, PP_ALIGN.LEFT)
    text(slide, "Both decisions are now upstream", 3.00, 5.78, 2.16, .25,
         10.5, RED, True, PP_ALIGN.CENTER)

    arrow(slide, 2.40, 2.44, 2.68, 2.75, DC)
    arrow(slide, 2.40, 5.25, 2.68, 4.93, CHEMBL)

    app_graph_card(slide, 5.84, 1.96)
    arrow(slide, 5.48, 3.86, 5.80, 3.86, RED)

    box(slide, 8.18, 2.05, 1.86, 3.66, DC_PALE, DC)
    text(slide, "ODIN ADAPTER", 8.35, 2.28, 1.52, .28,
         13.5, DC, True, PP_ALIGN.CENTER)
    text(slide, "DrugHarmonizer\nLigandNodeAdapter", 8.34, 2.78, 1.54, .65,
         11.3, INK, True, PP_ALIGN.CENTER)
    text(slide, "Creates a Ligand from\nselected export columns", 8.35, 3.72, 1.52, .60,
         10.4, INK, False, PP_ALIGN.CENTER)
    box(slide, 8.42, 4.60, 1.38, .72, WHITE, DC)
    text(slide, "framework source:\nDrug Harmonizer", 8.52, 4.72, 1.18, .48,
         9.5, DC, True, PP_ALIGN.CENTER)
    text(slide, "one adapter controls all included sources", 8.39, 5.39, 1.44, .18,
         8.4, RED, True, PP_ALIGN.CENTER)
    arrow(slide, 7.88, 3.86, 8.14, 3.86, MERGE)

    box(slide, 10.42, 1.64, 2.46, 4.72, WHITE, NAVY)
    text(slide, "CANONICAL DRUG", 10.61, 1.87, 2.08, .32,
         14.5, NAVY, True, PP_ALIGN.CENTER)
    text(slide, "IFXDrug:000042", 10.61, 2.27, 2.08, .28,
         12, RESOLVER, True, PP_ALIGN.CENTER)
    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(10.72), Inches(2.67), Inches(1.86), Inches(.025)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = LINE
    divider.line.fill.background()

    text(slide, "name", 10.63, 2.86, .52, .22, 10.3, MUTED, True)
    text(slide, "Example Drug", 11.18, 2.86, 1.30, .22, 10.3, INK)
    provenance_tag(slide, "Drug Harmonizer", 10.65, 3.16, 1.16, RED)

    text(slide, "isDrug", 10.63, 3.58, .52, .22, 10.3, MUTED, True)
    text(slide, "true", 11.18, 3.58, 1.30, .22, 10.3, INK)
    provenance_tag(slide, "Drug Harmonizer", 10.65, 3.88, 1.16, RED)

    box(slide, 10.67, 4.42, 1.96, 1.48, RED_PALE, RED)
    text(slide, "ODIN no longer knows", 10.82, 4.56, 1.66, .25,
         10.5, RED, True, PP_ALIGN.CENTER)
    text(slide,
         "which source set each field\n\nor which omitted source\nfields were available",
         10.83, 4.88, 1.64, .83, 9.5, INK, False, PP_ALIGN.CENTER)

    arrow(slide, 10.08, 3.86, 10.38, 3.86, RED)

    box(slide, .55, 6.63, 12.23, .55, RED_PALE, RED)
    text(slide,
         "ODIN cannot include or exclude DrugCentral and ChEMBL independently—they arrive as one coupled input.",
         .80, 6.74, 11.73, .30, 12.5, RED, True, PP_ALIGN.CENTER)
    text(slide,
         "Adding a new source field also requires changing the Harmonizer, rebuilding its export, and then rebuilding ODIN.",
         .70, 7.22, 11.93, .15, 8.7, MUTED, False, PP_ALIGN.CENTER)


def main():
    prs = Presentation(SOURCE)
    add_slide(prs)
    prs.core_properties.title = "ODIN adapter and resolver boundary examples"
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
