from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("adapter_resolver_visual.pptx")

NAVY = RGBColor(20, 42, 67)
INK = RGBColor(37, 45, 54)
MUTED = RGBColor(91, 103, 116)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(246, 248, 250)
LINE = RGBColor(195, 205, 214)

DC = RGBColor(39, 112, 174)
DC_PALE = RGBColor(231, 241, 249)
CHEMBL = RGBColor(121, 75, 156)
CHEMBL_PALE = RGBColor(242, 235, 247)
RESOLVER = RGBColor(21, 137, 128)
RESOLVER_PALE = RGBColor(226, 246, 242)
MERGE = RGBColor(221, 153, 42)
MERGE_PALE = RGBColor(253, 244, 224)


def text(slide, value, x, y, w, h, size=14, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.03)
    frame.margin_top = frame.margin_bottom = Inches(.01)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def box(slide, x, y, w, h, fill, border, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(1.25)
    if rounded:
        shape.adjustments[0] = .07
    return shape


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=2.4):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    shape.line.end_arrowhead = True
    return shape


def field(slide, label, value, x, y, w, color=INK, size=11.5):
    text(slide, label, x, y, .70, .24, size, MUTED, True)
    text(slide, value, x + .72, y, w - .72, .24, size, color)


def source_card(slide, title_value, fields, x, y, fill, accent):
    box(slide, x, y, 2.18, 1.88, fill, accent)
    text(slide, title_value, x + .16, y + .13, 1.86, .42, 14, accent, True, PP_ALIGN.CENTER)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + .25), Inches(y + .61), Inches(1.68), Inches(.025))
    rule.fill.solid()
    rule.fill.fore_color.rgb = accent
    rule.line.fill.background()
    for index, (label, value) in enumerate(fields):
        value_size = 9.4 if label == "ID" else 11.5
        field(slide, label, value, x + .16, y + .75 + index * .31, 1.86, accent, value_size)


def adapter_card(slide, title_value, x, y, fill, accent):
    box(slide, x, y, 1.46, 1.52, fill, accent)
    text(slide, title_value, x + .10, y + .16, 1.26, .50, 13, accent, True, PP_ALIGN.CENTER)
    text(slide, "maps source\ncolumns into\nLigand fields", x + .12, y + .73, 1.22, .58,
         11, INK, False, PP_ALIGN.CENTER)


def model_card(slide, values, x, y, fill, accent):
    box(slide, x, y, 2.05, 1.88, fill, accent)
    text(slide, "ODIN Ligand", x + .15, y + .13, 1.75, .34, 14, accent, True, PP_ALIGN.CENTER)
    for index, value in enumerate(values):
        text(slide, value, x + .14, y + .60 + index * .32, 1.77, .26, 10.7, INK)


def provenance_tag(slide, label, x, y, w, fill):
    shape = box(slide, x, y, w, .25, fill, fill)
    shape.adjustments[0] = .18
    text(slide, label, x + .04, y + .015, w - .08, .20, 8.5, WHITE, True, PP_ALIGN.CENTER)


def build_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    text(slide, "Two source records become one drug—without losing their audit trail",
         .52, .27, 12.28, .52, 25, NAVY, True)
    text(slide, "Adapters decide what each source contributes. The ID resolver determines that both records describe the same entity.",
         .54, .84, 12.10, .32, 13, MUTED)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.55), Inches(1.25), Inches(.92), Inches(.04))
    rule.fill.solid()
    rule.fill.fore_color.rgb = RESOLVER
    rule.line.fill.background()

    # Source A lane
    source_card(
        slide,
        "DATA SOURCE A · DrugCentral",
        [("ID", "DrugCentral:DC123"), ("name", "Example Drug"), ("status", "approved")],
        .45, 1.58, DC_PALE, DC,
    )
    adapter_card(slide, "DrugCentral\nAdapter", 2.96, 1.76, DC_PALE, DC)
    model_card(
        slide,
        ["id = DrugCentral:DC123", "name = Example Drug", "isDrug = true"],
        4.74, 1.58, DC_PALE, DC,
    )

    # Source B lane
    source_card(
        slide,
        "DATA SOURCE B · ChEMBL",
        [("ID", "ChEMBL:CHEMBL456"), ("name", "Example Drug"), ("phase", "4")],
        .45, 4.42, CHEMBL_PALE, CHEMBL,
    )
    adapter_card(slide, "ChEMBL\nAdapter", 2.96, 4.60, CHEMBL_PALE, CHEMBL)
    model_card(
        slide,
        ["id = ChEMBL:CHEMBL456", "name = Example Drug", "max_phase = 4"],
        4.74, 4.42, CHEMBL_PALE, CHEMBL,
    )

    # Flow through adapters
    arrow(slide, 2.67, 2.52, 2.92, 2.52, DC)
    arrow(slide, 4.44, 2.52, 4.70, 2.52, DC)
    arrow(slide, 2.67, 5.36, 2.92, 5.36, CHEMBL)
    arrow(slide, 4.44, 5.36, 4.70, 5.36, CHEMBL)

    # Resolver
    box(slide, 7.18, 2.20, 1.82, 3.37, RESOLVER_PALE, RESOLVER)
    text(slide, "ID RESOLVER", 7.36, 2.42, 1.46, .32, 15, RESOLVER, True, PP_ALIGN.CENTER)
    text(slide, "WHO?", 7.55, 2.82, 1.08, .30, 13, RESOLVER, True, PP_ALIGN.CENTER)
    text(slide, "DrugCentral:DC123", 7.35, 3.31, 1.48, .28, 10.3, DC, True, PP_ALIGN.CENTER)
    text(slide, "↓", 7.78, 3.64, .60, .28, 16, RESOLVER, True, PP_ALIGN.CENTER)
    box(slide, 7.40, 3.98, 1.38, .58, WHITE, RESOLVER)
    text(slide, "IFXDrug:000042", 7.48, 4.12, 1.22, .25, 10.5, RESOLVER, True, PP_ALIGN.CENTER)
    text(slide, "↑", 7.78, 4.63, .60, .28, 16, RESOLVER, True, PP_ALIGN.CENTER)
    text(slide, "ChEMBL:CHEMBL456", 7.31, 4.96, 1.56, .28, 10.1, CHEMBL, True, PP_ALIGN.CENTER)
    text(slide, "same preferred ID", 7.37, 5.28, 1.44, .18, 9.3, MUTED, False, PP_ALIGN.CENTER)

    arrow(slide, 6.82, 2.52, 7.14, 3.52, DC)
    arrow(slide, 6.82, 5.36, 7.14, 4.88, CHEMBL)

    # Merge step
    merge = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(9.25), Inches(3.43), Inches(.82), Inches(.95))
    merge.fill.solid()
    merge.fill.fore_color.rgb = MERGE_PALE
    merge.line.color.rgb = MERGE
    merge.line.width = Pt(1.4)
    text(slide, "MERGE", 9.34, 3.69, .64, .26, 10.5, MERGE, True, PP_ALIGN.CENTER)
    arrow(slide, 9.04, 3.91, 9.21, 3.91, RESOLVER)

    # Final canonical record and explicit audit trail
    box(slide, 10.30, 1.64, 2.58, 4.72, WHITE, NAVY)
    text(slide, "CANONICAL DRUG", 10.50, 1.86, 2.18, .32, 15, NAVY, True, PP_ALIGN.CENTER)
    text(slide, "IFXDrug:000042", 10.50, 2.25, 2.18, .28, 12.5, RESOLVER, True, PP_ALIGN.CENTER)
    rule2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.62), Inches(2.66), Inches(1.94), Inches(.025))
    rule2.fill.solid()
    rule2.fill.fore_color.rgb = LINE
    rule2.line.fill.background()

    text(slide, "name", 10.52, 2.84, .57, .24, 10.5, MUTED, True)
    text(slide, "Example Drug", 11.11, 2.84, 1.42, .24, 10.5, INK)
    provenance_tag(slide, "DrugCentral", 10.54, 3.15, .84, DC)
    provenance_tag(slide, "ChEMBL", 11.46, 3.15, .66, CHEMBL)

    text(slide, "isDrug", 10.52, 3.58, .57, .24, 10.5, MUTED, True)
    text(slide, "true", 11.11, 3.58, 1.42, .24, 10.5, INK)
    provenance_tag(slide, "DrugCentral", 10.54, 3.89, .84, DC)

    text(slide, "max phase", 10.52, 4.32, .73, .24, 10.5, MUTED, True)
    text(slide, "4", 11.28, 4.32, 1.25, .24, 10.5, INK)
    provenance_tag(slide, "ChEMBL", 10.54, 4.63, .66, CHEMBL)

    text(slide, "equivalent IDs", 10.52, 5.08, 1.08, .24, 10.5, MUTED, True)
    text(slide, "DC123 · CHEMBL456", 10.52, 5.38, 1.98, .24, 10.2, INK)
    provenance_tag(slide, "resolver", 10.54, 5.70, .66, RESOLVER)

    arrow(slide, 10.10, 3.91, 10.26, 3.91, MERGE)

    text(slide, "WHAT", 3.22, 3.48, .95, .24, 10.5, MUTED, True, PP_ALIGN.CENTER)
    text(slide, "WHAT", 3.22, 6.30, .95, .24, 10.5, MUTED, True, PP_ALIGN.CENTER)
    text(slide, "WHO", 7.62, 5.75, .95, .24, 10.5, MUTED, True, PP_ALIGN.CENTER)

    text(slide, "Each source remains independently selectable and auditable.",
         .55, 6.83, 12.22, .28, 13, NAVY, True, PP_ALIGN.CENTER)
    text(slide, "Illustrative identifiers", .57, 7.18, 2.0, .15, 8.5, MUTED)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "How ODIN adapters, resolvers, and merging work"
    prs.core_properties.author = "IFX_ODIN"
    build_slide(prs)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
