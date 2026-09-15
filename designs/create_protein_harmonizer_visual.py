"""Append a protein_ids.tsv example to the user-edited adapter/resolver slide."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from create_adapter_resolver_visual import (
    CHEMBL,
    CHEMBL_PALE,
    INK,
    LINE,
    MERGE,
    MERGE_PALE,
    MUTED,
    NAVY,
    RESOLVER,
    RESOLVER_PALE,
    WHITE,
    adapter_card,
    arrow,
    box,
    field,
    model_card,
    provenance_tag,
    text,
)


HERE = Path(__file__).parent
SOURCE = HERE / "adapter_resolver_visual.pptx"
OUT = HERE / "adapter_resolver_visual_with_protein.pptx"

HARMONIZER = CHEMBL
HARMONIZER_PALE = CHEMBL_PALE
UNIPROT = RESOLVER
UNIPROT_PALE = RESOLVER_PALE


def protein_source_card(slide, x, y):
    box(slide, x, y, 2.18, 2.02, HARMONIZER_PALE, HARMONIZER)
    text(slide, "TARGET HARMONIZER", x + .16, y + .10, 1.86, .25,
         11, HARMONIZER, True, PP_ALIGN.CENTER)
    text(slide, "protein_ids.tsv", x + .16, y + .37, 1.86, .31,
         15, HARMONIZER, True, PP_ALIGN.CENTER)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x + .25), Inches(y + .75), Inches(1.68), Inches(.025)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = HARMONIZER
    line.line.fill.background()
    field(slide, "ID", "IFXProtein:000042", x + .16, y + .88, 1.86, HARMONIZER, 9.4)
    field(slide, "UniProt", "P12345", x + .16, y + 1.18, 1.86, HARMONIZER, 10.5)
    field(slide, "score", "0.98", x + .16, y + 1.48, 1.86, HARMONIZER, 10.5)
    field(slide, "method", "consensus", x + .16, y + 1.78, 1.86, HARMONIZER, 10.5)


def uniprot_source_card(slide, x, y):
    box(slide, x, y, 2.18, 2.02, UNIPROT_PALE, UNIPROT)
    text(slide, "PRIMARY SOURCE · UniProt", x + .14, y + .14, 1.90, .38,
         13, UNIPROT, True, PP_ALIGN.CENTER)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x + .25), Inches(y + .62), Inches(1.68), Inches(.025)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = UNIPROT
    line.line.fill.background()
    field(slide, "accession", "P12345", x + .16, y + .78, 1.86, UNIPROT, 10.5)
    field(slide, "name", "Example receptor", x + .16, y + 1.08, 1.86, UNIPROT, 10.5)
    field(slide, "function", "Signal receptor", x + .16, y + 1.38, 1.86, UNIPROT, 10.5)
    field(slide, "sequence", "M…", x + .16, y + 1.68, 1.86, UNIPROT, 10.5)


def protein_model_card(slide, values, x, y, fill, accent):
    box(slide, x, y, 2.05, 2.02, fill, accent)
    text(slide, "ODIN Protein", x + .15, y + .13, 1.75, .34,
         14, accent, True, PP_ALIGN.CENTER)
    for index, value in enumerate(values):
        text(slide, value, x + .14, y + .59 + index * .31, 1.77, .25, 10.2, INK)


def add_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    text(slide, "One harmonizer output supports both identity and enrichment",
         .52, .27, 12.28, .52, 25, NAVY, True)
    text(slide,
         "protein_ids.tsv is used twice: its mappings build the resolver, while its scores and methods enrich Protein records.",
         .54, .84, 12.10, .32, 13, MUTED)
    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(.55), Inches(1.25), Inches(.92), Inches(.04)
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = RESOLVER
    rule.line.fill.background()

    protein_source_card(slide, .45, 1.54)
    adapter_card(slide, "Target Graph\nProtein Adapter", 2.96, 1.79, HARMONIZER_PALE, HARMONIZER)
    protein_model_card(
        slide,
        [
            "id = IFXProtein:000042",
            "protein_name_score = 0.98",
            "protein_name_method = consensus",
            "mapping_ratio = 1.0",
        ],
        4.74, 1.54, HARMONIZER_PALE, HARMONIZER,
    )

    uniprot_source_card(slide, .45, 4.39)
    adapter_card(slide, "UniProt\nProtein Adapter", 2.96, 4.64, UNIPROT_PALE, UNIPROT)
    protein_model_card(
        slide,
        [
            "id = UniProtKB:P12345",
            "name = Example receptor",
            "function = Signal receptor",
            "sequence = M…",
        ],
        4.74, 4.39, UNIPROT_PALE, UNIPROT,
    )

    arrow(slide, 2.67, 2.55, 2.92, 2.55, HARMONIZER)
    arrow(slide, 4.44, 2.55, 4.70, 2.55, HARMONIZER)
    arrow(slide, 2.67, 5.40, 2.92, 5.40, UNIPROT)
    arrow(slide, 4.44, 5.40, 4.70, 5.40, UNIPROT)

    box(slide, 7.18, 2.18, 1.82, 3.42, RESOLVER_PALE, RESOLVER)
    text(slide, "TARGET GRAPH", 7.36, 2.38, 1.46, .25,
         10.5, RESOLVER, True, PP_ALIGN.CENTER)
    text(slide, "PROTEIN RESOLVER", 7.31, 2.66, 1.56, .47,
         14, RESOLVER, True, PP_ALIGN.CENTER)
    text(slide, "WHO?", 7.55, 3.17, 1.08, .25,
         12.5, RESOLVER, True, PP_ALIGN.CENTER)
    text(slide, "UniProtKB:P12345", 7.33, 3.66, 1.52, .26,
         10.0, UNIPROT, True, PP_ALIGN.CENTER)
    text(slide, "↓", 7.78, 4.00, .60, .27,
         16, RESOLVER, True, PP_ALIGN.CENTER)
    box(slide, 7.40, 4.34, 1.38, .58, WHITE, RESOLVER)
    text(slide, "IFXProtein:000042", 7.45, 4.48, 1.28, .25,
         9.8, RESOLVER, True, PP_ALIGN.CENTER)
    text(slide, "preferred Protein ID", 7.37, 5.12, 1.44, .20,
         9.2, MUTED, False, PP_ALIGN.CENTER)

    # Normal adapter output passes through the resolver.
    arrow(slide, 6.82, 2.55, 7.14, 3.60, HARMONIZER)
    arrow(slide, 6.82, 5.40, 7.14, 4.88, UNIPROT)

    # The same protein_ids.tsv also supplies the resolver's lookup index.
    direct_color = RESOLVER
    vertical1 = slide.shapes.add_connector(
        1, Inches(1.88), Inches(3.58), Inches(1.88), Inches(3.92)
    )
    vertical1.line.color.rgb = direct_color
    vertical1.line.width = Pt(2.0)
    horizontal = slide.shapes.add_connector(
        1, Inches(1.88), Inches(3.92), Inches(7.02), Inches(3.92)
    )
    horizontal.line.color.rgb = direct_color
    horizontal.line.width = Pt(2.0)
    final_leg = slide.shapes.add_connector(
        1, Inches(7.02), Inches(3.92), Inches(7.14), Inches(3.92)
    )
    final_leg.line.color.rgb = direct_color
    final_leg.line.width = Pt(2.0)
    final_leg.line.end_arrowhead = True
    text(slide, "the same TSV builds the resolver index",
         2.55, 3.65, 3.45, .24, 10.3, RESOLVER, True, PP_ALIGN.CENTER)

    merge = slide.shapes.add_shape(
        MSO_SHAPE.HEXAGON, Inches(9.25), Inches(3.45), Inches(.82), Inches(.95)
    )
    merge.fill.solid()
    merge.fill.fore_color.rgb = MERGE_PALE
    merge.line.color.rgb = MERGE
    merge.line.width = Pt(1.4)
    text(slide, "MERGE", 9.34, 3.71, .64, .26,
         10.5, MERGE, True, PP_ALIGN.CENTER)
    arrow(slide, 9.04, 3.93, 9.21, 3.93, RESOLVER)

    box(slide, 10.30, 1.56, 2.58, 4.92, WHITE, NAVY)
    text(slide, "CANONICAL PROTEIN", 10.48, 1.77, 2.22, .34,
         14.5, NAVY, True, PP_ALIGN.CENTER)
    text(slide, "IFXProtein:000042", 10.48, 2.17, 2.22, .28,
         12, RESOLVER, True, PP_ALIGN.CENTER)
    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(10.62), Inches(2.57), Inches(1.94), Inches(.025)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = LINE
    divider.line.fill.background()

    text(slide, "name score", 10.52, 2.72, .82, .22, 10.2, MUTED, True)
    text(slide, "0.98", 11.36, 2.72, 1.15, .22, 10.2, INK)
    provenance_tag(slide, "Target Harmonizer", 10.54, 3.00, 1.20, HARMONIZER)

    text(slide, "name method", 10.52, 3.36, .92, .22, 10.2, MUTED, True)
    text(slide, "consensus", 11.46, 3.36, 1.05, .22, 10.2, INK)
    provenance_tag(slide, "Target Harmonizer", 10.54, 3.64, 1.20, HARMONIZER)

    text(slide, "mapping ratio", 10.52, 4.00, .98, .22, 10.2, MUTED, True)
    text(slide, "1.0", 11.52, 4.00, .99, .22, 10.2, INK)
    provenance_tag(slide, "Target Harmonizer", 10.54, 4.28, 1.20, HARMONIZER)

    text(slide, "function", 10.52, 4.64, .64, .22, 10.2, MUTED, True)
    text(slide, "Signal receptor", 11.18, 4.64, 1.33, .22, 10.2, INK)
    provenance_tag(slide, "UniProt", 10.54, 4.92, .66, UNIPROT)

    text(slide, "sequence", 10.52, 5.28, .78, .22, 10.0, MUTED, True)
    text(slide, "M…", 11.34, 5.28, 1.17, .22, 10.2, INK)
    provenance_tag(slide, "UniProt", 10.54, 5.56, .66, UNIPROT)

    text(slide, "equivalent ID", 10.52, 5.92, .98, .22, 9.4, MUTED, True)
    text(slide, "UniProtKB:P12345", 11.52, 5.92, .97, .22, 8.5, INK)
    provenance_tag(slide, "resolver", 10.54, 6.18, .66, RESOLVER)

    arrow(slide, 10.10, 3.93, 10.26, 3.93, MERGE)
    text(slide, "WHAT", 3.22, 3.47, .95, .24,
         10.5, MUTED, True, PP_ALIGN.CENTER)
    text(slide, "WHAT", 3.22, 6.31, .95, .24,
         10.5, MUTED, True, PP_ALIGN.CENTER)
    text(slide, "WHO", 7.62, 5.80, .95, .24,
         10.5, MUTED, True, PP_ALIGN.CENTER)

    text(slide,
         "The harmonizer adds its own value while UniProt remains independently available and auditable.",
         .55, 6.88, 12.22, .28, 13, NAVY, True, PP_ALIGN.CENTER)
    text(slide, "Illustrative values", .57, 7.20, 2.0, .15, 8.5, MUTED)


def main():
    prs = Presentation(SOURCE)
    add_slide(prs)
    prs.core_properties.title = "ODIN adapter and resolver examples"
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
